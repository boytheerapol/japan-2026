#!/usr/bin/env python3
"""Convert the trip Excel workbook into site/data/trip.json.

Usage:
    python3 scripts/build_data.py [path/to/itinerary.xlsx]

The workbook itself is NOT committed (it holds booking numbers and real
spend). Only the sanitized JSON produced here is published. The script
checks sheet names and header rows first and exits non-zero if the layout
changed, so a reshaped workbook never silently publishes wrong data.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "source" / "itinerary.xlsx"
META_JSON = ROOT / "source" / "day_meta.json"
OUT_JSON = ROOT / "site" / "data" / "trip.json"

TRIP_YEAR = 2026
THAI_MONTHS = {"ต.ค.": 10, "พ.ย.": 11}
THAI_WEEKDAYS = {  # Python weekday(): Monday = 0
    "จันทร์": 0, "อังคาร": 1, "พุธ": 2, "พฤหัส": 3, "พฤหัสบดี": 3,
    "ศุกร์": 4, "เสาร์": 5, "อาทิตย์": 6,
}
WEEKDAY_TH = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]

# Expected header row (row 4) per sheet: {column letter: header text}.
EXPECTED_HEADERS = {
    "Itinerary Plan": {"B": "วันที่", "C": "เส้นทางหลัก", "D": "ออก", "E": "ระยะขับ",
                       "F": "ตะวันขึ้น", "G": "ตะวันตก", "H": "ที่พักคืนนั้น"},
    "Daily Checklist": {"C": "เวลา", "D": "ประเภท", "E": "★", "F": "กิจกรรม",
                        "G": "เวลาเปิด–ปิด", "H": "วันปิด / ข้อจำกัด",
                        "I": "รายละเอียด / เคล็ดลับ", "J": "ใช้เวลา / ระยะทาง",
                        "K": "ต้องจอง", "L": "¥/คน", "M": "verify"},
    "🗺️ Route Links": {"A": "วัน", "B": "เส้นทาง", "C": "วิธีเดินทาง",
                       "D": "จุดแวะ (ตามลำดับ)", "E": "จุด", "F": "ระยะ / เวลา",
                       "G": "เปิด Google Maps", "H": "หมายเหตุ"},
    "🅿️ Parking": {"A": "วัน", "B": "สถานที่", "C": "🅿️ ที่จอดหลัก", "D": "ค่าจอด",
                   "E": "ข้อจำกัด / จำนวนคัน", "F": "เดินถึงจุด", "G": "เปิดลานหลัก",
                   "H": "🅿️ ที่จอดสำรอง", "I": "เปิดลานสำรอง", "J": "โทร"},
    "Fuji Viewpoints": {"A": "โซน", "B": "★", "C": "จุดชมวิว", "D": "ลักษณะวิว",
                        "E": "ทิศที่เห็นฟูจิ + ช่วงแสงที่ดีที่สุด", "F": "การเข้าถึง",
                        "G": "ค่าเข้า / ที่จอด", "H": "อยู่ในแผนวันไหน"},
    "Food & Souvenirs": {"A": "โซน / เมือง", "B": "หมวด", "C": "★", "D": "รายการ",
                         "E": "ร้าน / ที่ซื้อ", "F": "วันในแผน", "G": "รายละเอียด",
                         "H": "ราคาโดยประมาณ"},
    "Tokyo by Area": {"A": "ย่าน", "B": "หมวด", "C": "ชื่อสถานที่ / ร้าน",
                      "D": "จุดเด่น / เมนูแนะนำ", "E": "ค่าเข้า / ราคา",
                      "F": "เวลาเปิด-ปิด / วันปิด", "G": "สถานี / สายที่ลง", "H": "ข้อมูล"},
    "Pre-Trip": {"A": "ช่วงเวลา", "B": "สิ่งที่ต้องทำ", "C": "รายละเอียด / เหตุผล",
                 "D": "ความสำคัญ"},
}

# The workbook is written for the driver ("คุณ" = the trip owner). Friends read
# the site too, so rewrite second-person phrasing into neutral wording.
SANITIZE = [
    (r"\s*\(?ลิงก์เดิมของคุณชี้ลานนี้\)?", ""),
    (r"\s*ตามลิงก์เดิมของคุณ", ""),
    (r"\s*\(?ลานที่คุณตั้งใจจอด\s*=\s*", " ("),
    (r"\s*\(?ลิงก์ที่คุณเช็ค\)?", ""),
    (r"\s*ตามลิงก์ที่คุณเช็ค", ""),
    (r"\s*ตามที่คุณเช็ค", ""),
    (r"\s*เพิ่มตามลิงก์ของคุณ", ""),
    (r"\s*ตามลิงก์ของคุณ", ""),
    (r"ปรับจากลิงก์ของคุณให้", ""),
    (r"ลิงก์ของคุณ\s*·?\s*", ""),
    (r"\s*\+\s*ลิงก์ของคุณ", ""),
    (r"🅰️ คุณ \+ คู่ พัก", "🅰️ กลุ่ม A พัก"),
    (r"🅱️ เพื่อนพัก", "🅱️ กลุ่ม B พัก"),
    (r"🅱️ เพื่อนแยกไป", "🅱️ กลุ่ม B แยกไป"),
    (r"\(คุณ \+ คู่\)", "(กลุ่ม A)"),
    (r"\(เพื่อน\)", "(กลุ่ม B)"),
    (r"คุณขับต่อคนเดียว", "คนขับขับต่อคนเดียว"),
    (r"คุณขับคนเดียว", "คนขับขับคนเดียว"),
    (r"\(คุณขับคนเดียว\)", "(คนขับคนเดียว)"),
    (r"\(คุณคนเดียว\)", "(คนขับคนเดียว)"),
    (r"แล้วคุณขับ", "แล้วคนขับขับ"),
    (r"คุณส่ง", "คนขับส่ง"),
    (r"\s*[•·]?\s*ยอดจ(?:อง|่าย)จริง\s*¥[\d,]+[^·•\n]*", ""),
    (r"⚠️ แก้จาก v\d+:\s*", "⚠️ "),
    (r"ซึ่งคือคุณ", ""),
    (r"ซึ่งคือเพื่อน", ""),
    (r"\(\s*\)", ""),
    (r"[ \t]{2,}", " "),
]


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").strip()
    for pattern, repl in SANITIZE:
        text = re.sub(pattern, repl, text)
    return text.strip()


def link(cell) -> str:
    return cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else ""


def fail(msg: str) -> None:
    sys.exit(f"[build_data] ERROR: {msg}")


def check_layout(wb) -> None:
    for sheet, headers in EXPECTED_HEADERS.items():
        if sheet not in wb.sheetnames:
            fail(f"sheet '{sheet}' not found. Sheets: {wb.sheetnames}")
        ws = wb[sheet]
        for col, expected in headers.items():
            got = clean(ws[f"{col}4"].value)
            if got != expected:
                fail(f"sheet '{sheet}' header {col}4 is '{got}', expected '{expected}'")


def day_numbers(label: str) -> list[int]:
    """'Day 1' -> [1]; 'Day 1–3' -> [1, 2, 3]; 'Day 8 (สำรอง)' -> [8]."""
    m = re.match(r"Day\s*(\d+)(?:\s*[–-]\s*(\d+))?", label or "")
    if not m:
        return []
    a = int(m.group(1))
    b = int(m.group(2) or a)
    return list(range(a, b + 1))


def parse_date(label: str) -> dt.date:
    m = re.match(r"(\d+)\s*(ต\.ค\.|พ\.ย\.)", label)
    if not m:
        fail(f"cannot parse date '{label}'")
    return dt.date(TRIP_YEAR, THAI_MONTHS[m.group(2)], int(m.group(1)))


def yen(value) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def closed_conflict(text: str, date: dt.date) -> str:
    """Return a warning if `text` says the place closes on this date's weekday."""
    for m in re.finditer(r"ปิด(?:วัน)?(จันทร์|อังคาร|พุธ|พฤหัสบดี|พฤหัส|ศุกร์|เสาร์|อาทิตย์)"
                         r"(?:ที่\s*([\d\s,และ]+))?", text):
        wd = THAI_WEEKDAYS[m.group(1)]
        if wd != date.weekday():
            continue
        nths = [int(n) for n in re.findall(r"\d", m.group(2) or "")]
        nth_of_month = (date.day - 1) // 7 + 1
        if nths and nth_of_month not in nths:
            continue
        # The sheet often already notes "X เป็นพุธ เปิดปกติ" — that is not a conflict.
        if "เปิดปกติ" in text:
            continue
        return f"ปิดวัน{WEEKDAY_TH[wd]} — ตรงกับวันนี้"
    return ""


def split_detail(text: str) -> dict:
    """Split the multi-line detail cell into tagged parts."""
    out = {"detail": [], "parking": [], "backup": [], "rule": []}
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("🅿️"):
            out["parking"].append(line.removeprefix("🅿️").strip())
        elif line.startswith("🔁"):
            out["backup"].append(line.removeprefix("🔁").strip())
        elif line.startswith("⏱"):
            out["rule"].append(line.removeprefix("⏱").strip())
        else:
            # Inline markers inside a paragraph (e.g. "... · 🔁 ...").
            parts = re.split(r"\s*(?=🅿️|🔁)", line)
            out["detail"].append(parts[0])
            for p in parts[1:]:
                key = "parking" if p.startswith("🅿️") else "backup"
                out[key].append(p[2:].strip() if key == "parking" else p[1:].strip())
    return {k: v for k, v in out.items() if v}


def read_overview(wb) -> dict[int, dict]:
    ws = wb["Itinerary Plan"]
    days = {}
    for row in ws.iter_rows(min_row=5):
        label = clean(row[0].value)
        nums = day_numbers(label)
        if len(nums) != 1 or not row[1].value:
            continue
        n = nums[0]
        date = parse_date(clean(row[1].value))
        days[n] = {
            "n": n,
            "date": date.isoformat(),
            "dateLabel": clean(row[1].value),
            "weekday": WEEKDAY_TH[date.weekday()],
            "route": clean(row[2].value),
            "depart": clean(row[3].value),
            "drive": clean(row[4].value),
            "sunrise": clean(row[5].value),
            "sunset": clean(row[6].value),
            "hotel": clean(row[7].value),
            "items": [], "alerts": [], "notes": [], "routes": [], "parking": [],
        }
    if len(days) != 11:
        fail(f"expected 11 days in 'Itinerary Plan', found {len(days)}")
    return days


def read_checklist(wb, days: dict) -> None:
    ws = wb["Daily Checklist"]
    for row in ws.iter_rows(min_row=5):
        nums = day_numbers(clean(row[0].value))
        if len(nums) != 1 or not row[1].value:  # section header rows have no date
            continue
        day = days[nums[0]]
        date = dt.date.fromisoformat(day["date"])
        kind = clean(row[3].value)
        item = {
            "time": clean(row[2].value),
            "type": kind,
            "stars": clean(row[4].value).count("★"),
            "name": clean(row[5].value),
            "hours": clean(row[6].value),
            "limits": clean(row[7].value),
            **split_detail(clean(row[8].value)),
            "duration": clean(row[9].value),
            "booking": clean(row[10].value),
            "yen": yen(row[11].value),
            "verified": clean(row[12].value) if isinstance(row[12].value, str) else
            (row[12].value.date().isoformat() if row[12].value else ""),
        }
        warn = closed_conflict(f"{item['hours']} {item['limits']}", date)
        if warn:
            item["closedWarning"] = warn
        item = {k: v for k, v in item.items() if v not in ("", [], 0) or k == "yen"}
        if kind.startswith("⚠️"):
            day["alerts"].append(item)
        elif kind.startswith("ℹ️"):
            day["notes"].append(item)
        else:
            day["items"].append(item)


def read_routes(wb, days: dict) -> None:
    ws = wb["🗺️ Route Links"]
    for row in ws.iter_rows(min_row=5):
        label = clean(row[0].value)
        nums = day_numbers(label)
        url = link(row[6])
        if len(nums) != 1 or not url:
            continue
        mode = clean(row[2].value)
        days[nums[0]]["routes"].append({
            "label": re.sub(r"^Day\s*\d+\s*", "", label).strip() or "ทั้งวัน",
            "title": clean(row[1].value),
            "mode": "drive" if "ขับ" in mode else "walk" if "เดิน" in mode else "transit",
            "modeLabel": mode,
            "stops": [s.strip() for s in clean(row[3].value).split("→") if s.strip()],
            "distance": clean(row[5].value),
            "url": url,
            "note": clean(row[7].value),
            "backup": "🔁" in label,
        })


def read_parking(wb, days: dict) -> list[dict]:
    ws = wb["🅿️ Parking"]
    all_lots = []
    for row in ws.iter_rows(min_row=5):
        label = clean(row[0].value)
        nums = day_numbers(label)
        if not nums or not row[1].value:
            continue
        lot = {
            "id": f"p{row[0].row}",
            "days": nums,
            "dayLabel": label,
            "place": clean(row[1].value),
            "lot": clean(row[2].value),
            "fee": clean(row[3].value),
            "limits": clean(row[4].value),
            "walk": clean(row[5].value),
            "url": link(row[6]),
            "alt": clean(row[7].value),
            "altUrl": link(row[8]),
            "phone": clean(row[9].value),
            "backupPlan": "🔁" in label or "สำรอง" in label,
        }
        lot = {k: v for k, v in lot.items() if v not in ("", None)}
        all_lots.append(lot)
        for n in nums:
            days[n]["parking"].append(lot["id"])
    return all_lots


def read_table(wb, sheet: str, cols: dict[str, str], group_col: str = "A") -> list[dict]:
    """Generic reader: rows with only column A filled are group headers."""
    ws = wb[sheet]
    groups: list[dict] = []
    for row in ws.iter_rows(min_row=5):
        cells = {openpyxl.utils.get_column_letter(c.column): c for c in row}
        values = {k: clean(cells[k].value) if k in cells else "" for k in cols.values()}
        first = clean(cells["A"].value) if "A" in cells else ""
        others = [clean(c.value) for c in row[1:] if c.value is not None]
        if first and not others:
            groups.append({"group": first, "rows": []})
            continue
        if not any(values.values()):
            continue
        if not groups:
            groups.append({"group": "", "rows": []})
        entry = {name: values[col] for name, col in cols.items() if values[col]}
        if "stars" in entry:
            entry["stars"] = entry["stars"].count("★")
        groups[-1]["rows"].append(entry)
    return [g for g in groups if g["rows"]]


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        fail(f"workbook not found: {xlsx}")
    wb = openpyxl.load_workbook(xlsx)
    check_layout(wb)
    meta = json.loads(META_JSON.read_text(encoding="utf-8"))

    days = read_overview(wb)
    read_checklist(wb, days)
    read_routes(wb, days)
    parking = read_parking(wb, days)
    for n, day in days.items():
        extra = meta["days"].get(str(n), {})
        day.update({k: extra[k] for k in ("title", "subtitle", "intro") if k in extra})
        day["budgetYen"] = sum(i.get("yen", 0) for i in day["items"])
        if not day["items"]:
            fail(f"Day {n} has no checklist items")

    guides = {
        "food": read_table(wb, "Food & Souvenirs", {
            "category": "B", "stars": "C", "name": "D", "shop": "E", "day": "F",
            "detail": "G", "price": "H"}),
        "tokyo": read_table(wb, "Tokyo by Area", {
            "category": "B", "name": "C", "detail": "D", "price": "E", "hours": "F",
            "station": "G", "info": "H"}),
        "fuji": read_table(wb, "Fuji Viewpoints", {
            "stars": "B", "name": "C", "view": "D", "light": "E", "access": "F",
            "fee": "G", "day": "H", "distance": "I", "note": "J"}),
        "pretrip": read_table(wb, "Pre-Trip", {
            "task": "B", "detail": "C", "priority": "D"}),
    }
    for g in guides["tokyo"]:  # Tokyo sheet embeds stars in the name ("★★ Tsukiji ...")
        for r in g["rows"]:
            m = re.match(r"^(★+)\s*", r.get("name", ""))
            if m:
                r["stars"] = len(m.group(1))
                r["name"] = r["name"][m.end():]

    trip = {
        "title": "Japan 2026",
        "subtitle": clean(wb["ปก"]["A3"].value),
        "tagline": meta.get("tagline", ""),
        "overview": meta.get("overview", ""),
        "start": days[1]["date"],
        "end": days[11]["date"],
        "firstFlight": "2026-10-29T01:40:00+07:00",
        "packing": meta.get("packing", []),
        "days": [days[n] for n in sorted(days)],
        "parking": parking,
        "guides": guides,
        "generated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "source": xlsx.name,
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(trip, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    leftovers = sorted({m for m in re.findall(r".{0,20}คุณ.{0,20}", OUT_JSON.read_text("utf-8"))})
    print(f"[build_data] wrote {OUT_JSON.relative_to(ROOT)}: "
          f"{sum(len(d['items']) for d in trip['days'])} items, "
          f"{sum(len(d['routes']) for d in trip['days'])} routes, {len(parking)} parking lots")
    if leftovers:
        print("[build_data] note: second-person text still present:")
        for s in leftovers:
            print("   ", s.replace("\n", " "))


if __name__ == "__main__":
    main()
